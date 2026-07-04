from __future__ import annotations

import shutil
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from config import OUTPUT_COLUMNS, STATUS_INVALID_FSN, build_url
from models import ScrapeResult


def normalize_fsn(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return text.upper()


def read_input_excel(input_path: str | Path, fsn_column: str) -> pd.DataFrame:
    path = Path(input_path)
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")

    df = pd.read_excel(path, dtype=str).fillna("")
    df.columns = [str(column).strip() for column in df.columns]
    if fsn_column not in df.columns:
        raise ValueError(f"'{fsn_column}' column not found. Available columns: {list(df.columns)}")
    df[fsn_column] = df[fsn_column].map(normalize_fsn)
    return df


def unique_fsns(df: pd.DataFrame, fsn_column: str, limit: int | None = None) -> list[str]:
    values = [normalize_fsn(value) for value in df[fsn_column].tolist()]
    unique = list(dict.fromkeys(fsn for fsn in values if fsn))
    return unique[:limit] if limit else unique


def is_valid_fsn(fsn: str) -> bool:
    if not fsn:
        return False
    return fsn.isalnum() and 8 <= len(fsn) <= 30


def merge_results_into_original_excel(
    input_df: pd.DataFrame,
    results: dict[str, ScrapeResult],
    fsn_column: str = "FSN",
) -> pd.DataFrame:
    output_df = input_df.copy()
    for column in OUTPUT_COLUMNS:
        if column not in output_df.columns:
            output_df[column] = ""

    for index, row in output_df.iterrows():
        fsn = normalize_fsn(row.get(fsn_column, ""))
        if not fsn:
            for column in OUTPUT_COLUMNS:
                output_df.at[index, column] = ""
            continue

        result = results.get(fsn)
        if result is None:
            continue
        for column, value in result.to_output_values().items():
            output_df.at[index, column] = value

    original_columns = list(input_df.columns)
    appended_columns = [column for column in OUTPUT_COLUMNS if column not in original_columns]
    existing_output_columns = [column for column in OUTPUT_COLUMNS if column in original_columns]
    ordered_columns = original_columns + appended_columns
    for column in existing_output_columns:
        if column not in ordered_columns:
            ordered_columns.append(column)
    return output_df[ordered_columns]


def _sheet_headers(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for column in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=column).value
        if value is not None:
            headers[str(value).strip()] = column
    return headers


def write_results_to_workbook(
    input_path: str | Path,
    output_path: str | Path,
    results: dict[str, ScrapeResult],
    fsn_column: str,
) -> None:
    input_path = Path(input_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if input_path.resolve() != output_path.resolve():
        shutil.copy2(input_path, output_path)

    wb = load_workbook(output_path)
    ws = wb.active
    headers = _sheet_headers(ws)
    if fsn_column not in headers:
        raise ValueError(f"'{fsn_column}' column not found in workbook header row")

    for column in OUTPUT_COLUMNS:
        if column not in headers:
            next_col = ws.max_column + 1
            ws.cell(row=1, column=next_col, value=column)
            headers[column] = next_col

    fsn_col = headers[fsn_column]
    for row in range(2, ws.max_row + 1):
        fsn = normalize_fsn(ws.cell(row=row, column=fsn_col).value)
        if not fsn:
            for column in OUTPUT_COLUMNS:
                ws.cell(row=row, column=headers[column], value="")
            continue
        result = results.get(fsn)
        if result is None:
            continue
        for column, value in result.to_output_values().items():
            ws.cell(row=row, column=headers[column], value=value)

    for column in OUTPUT_COLUMNS:
        col_letter = get_column_letter(headers[column])
        ws.column_dimensions[col_letter].width = max(ws.column_dimensions[col_letter].width or 0, 18)

    wb.save(output_path)


def load_resume_results(output_path: str | Path, fsn_column: str) -> dict[str, ScrapeResult]:
    path = Path(output_path)
    if not path.exists():
        return {}

    df = pd.read_excel(path, dtype=str).fillna("")
    df.columns = [str(column).strip() for column in df.columns]
    required = {fsn_column, *OUTPUT_COLUMNS}
    if not required.issubset(set(df.columns)):
        return {}

    resumed: dict[str, ScrapeResult] = {}
    for _, row in df.iterrows():
        fsn = normalize_fsn(row.get(fsn_column, ""))
        status = str(row.get("Scrape_Status", "")).strip()
        scraped_at = str(row.get("Scraped_At", "")).strip()
        if not fsn or not status or not scraped_at:
            continue
        if fsn in resumed:
            continue
        resumed[fsn] = ScrapeResult(
            fsn=fsn,
            url=str(row.get("Flipkart_URL", "")).strip() or build_url(fsn),
            deal_tag_found=str(row.get("Deal_Tag_Found", "")).strip().upper() == "TRUE",
            deal_tag=str(row.get("Deal_Tag", "")).strip(),
            discount_percentage=str(row.get("Discount_Percentage", "")).strip(),
            old_price=str(row.get("Old_Price", "")).strip(),
            new_price=str(row.get("New_Price", "")).strip(),
            status=status,
            reason=str(row.get("Scrape_Reason", "")).strip(),
            scraped_at=scraped_at,
            debug_path="",
        )
    return resumed


def invalid_fsn_results(fsns: Iterable[str], scraped_at: str) -> dict[str, ScrapeResult]:
    return {
        fsn: ScrapeResult.blank(
            fsn=fsn,
            status=STATUS_INVALID_FSN,
            reason="invalid FSN format",
            scraped_at=scraped_at,
        )
        for fsn in fsns
        if fsn and not is_valid_fsn(fsn)
    }
